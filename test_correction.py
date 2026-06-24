"""
AI correction smoke-test — loads template + placeholders from example_replacement.json,
runs GPT + Claude in parallel, prints before/after diff with timing.

Usage:
    python test_correction.py
    python test_correction.py path/to/other.json
"""

import base64
import io
import json
import os
import sys
import time

sys.path.insert(0, os.path.dirname(__file__))

# Force UTF-8 output on Windows
if sys.stdout.encoding and sys.stdout.encoding.lower() != "utf-8":
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

from docx import Document

from word_constructor.ai_correction.claude_checker import claude_available
from word_constructor.ai_correction.extraction import (
    document_full_text,
    extract_placeholder_occurrences,
)
from word_constructor.ai_correction.openai_client import request_ai_corrections
from word_constructor.ai_correction.rules import load_rules_config

JSON_PATH = sys.argv[1] if len(sys.argv) > 1 else "example_replacement.json"

# Keys that are not real placeholders (metadata mixed into the payload)
_SKIP_KEYS = {"МассивШапки", "ИспользоватьAI", "PromtAI", "UseAI"}


def load_payload(path: str):
    with open(path, encoding="utf-8") as f:
        data = json.load(f)

    # Decode embedded template
    b64 = data.get("content_base64", "")
    if not b64:
        raise ValueError("example_replacement.json has no content_base64 field")
    doc_bytes = base64.b64decode(b64)
    doc = Document(io.BytesIO(doc_bytes))

    # Filter to only real string placeholders
    raw = data.get("placeholders", {})
    slot_values = {
        k: str(v)
        for k, v in raw.items()
        if k not in _SKIP_KEYS and isinstance(v, str)
    }

    prompt_ai = str(raw.get("PromtAI") or data.get("PromtAI") or "")
    return doc, slot_values, prompt_ai


def run_gpt(full_text, occurrences, rules, slot_values, prompt_ai):
    return request_ai_corrections(
        full_text=full_text,
        occurrences=occurrences,
        rules=rules,
        prompt_ai=prompt_ai,
        placeholders=slot_values,
        log_key="test-run",
        timeout_seconds=60.0,
    )


def run_claude(full_text, slot_values, prompt_ai):
    from word_constructor.ai_correction.claude_checker_and_summarizer import (
        claude_correct_values,
    )
    return claude_correct_values(
        full_text=full_text,
        slot_values=slot_values,
        prompt_ai=prompt_ai,
        log_key="test-run",
        timeout_seconds=60.0,
    )


def main():
    print(f"\n{'='*65}")
    print(f"  JSON     : {JSON_PATH}")
    print(f"  GPT key  : {'SET' if os.environ.get('OPENAI_API_KEY') else 'MISSING !'}")
    print(f"  Claude   : {'available' if claude_available() else 'not configured !'}")
    print(f"{'='*65}\n")

    doc, slot_values, prompt_ai = load_payload(JSON_PATH)
    print(f"Placeholders ({len(slot_values)}):")
    for k, v in slot_values.items():
        print(f"  {k:<50} = {v!r}")

    print()
    t0 = time.perf_counter()

    full_text = document_full_text(doc)
    occurrences = extract_placeholder_occurrences(doc, slot_values)
    rules = load_rules_config()

    t_extract = time.perf_counter() - t0
    print(f"Extraction : {len(occurrences)} occurrences  ({t_extract*1000:.0f} ms)\n")

    import concurrent.futures

    gpt_per_key: dict[str, str] = {}
    claude_per_key: dict[str, str] = {}
    claude_summary = ""

    with concurrent.futures.ThreadPoolExecutor(max_workers=2) as pool:
        gpt_future = pool.submit(run_gpt, full_text, occurrences, rules, slot_values, prompt_ai)
        claude_future = pool.submit(run_claude, full_text, slot_values, prompt_ai) if claude_available() else None

        # GPT
        t_gpt = time.perf_counter()
        try:
            gpt_raw = gpt_future.result(timeout=70)
            for (k, _), v in gpt_raw.items():
                gpt_per_key[k] = v
            print(f"[GPT]    done  {(time.perf_counter()-t_gpt)*1000:.0f} ms  —  {len(gpt_per_key)} values")
        except Exception as exc:
            print(f"[GPT]    FAILED {(time.perf_counter()-t_gpt)*1000:.0f} ms  —  {type(exc).__name__}: {exc}")

        # Claude
        if claude_future:
            t_cl = time.perf_counter()
            try:
                cl_raw, claude_summary = claude_future.result(timeout=75)
                claude_per_key = cl_raw  # already dict[str, str]
                print(f"[Claude] done  {(time.perf_counter()-t_cl)*1000:.0f} ms  —  {len(claude_per_key)} values overridden")
            except Exception as exc:
                print(f"[Claude] FAILED {(time.perf_counter()-t_cl)*1000:.0f} ms  —  {type(exc).__name__}: {exc}")
        else:
            print("[Claude] skipped (not configured)")

    t_total = time.perf_counter() - t0

    # Merge: Claude wins
    final: dict[str, str] = {**slot_values, **gpt_per_key, **claude_per_key}

    # Report
    print(f"\n{'='*65}")
    print(f"  Total wall time: {t_total:.2f}s")
    print(f"{'='*65}")
    col_k = max(len(k) for k in slot_values) + 2
    col_b = 35
    header = f"  {'KEY':<{col_k}}  {'BEFORE':<{col_b}}  AFTER"
    print(f"\n{header}")
    print("-" * len(header))
    for k in slot_values:
        before = slot_values[k]
        after = final.get(k, before)
        src = ""
        if after != before:
            src = " (Claude)" if k in claude_per_key else " (GPT)"
        mark = "✓" if after != before else " "
        print(f"  {mark} {k:<{col_k}} {before:<{col_b}} {after}{src}")

    if claude_summary:
        print(f"\nClaude summary:\n  {claude_summary[:300]}")

    # Changes summary
    changed = [k for k in slot_values if final.get(k, slot_values[k]) != slot_values[k]]
    print(f"\n  Changed: {len(changed)}/{len(slot_values)} placeholders\n")


if __name__ == "__main__":
    main()
