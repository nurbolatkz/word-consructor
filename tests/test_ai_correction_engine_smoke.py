from __future__ import annotations

import json
from pathlib import Path
import unittest

from docx import Document

from word_constructor.ai_correction import (
    CorrectionEngine,
    KazakhAwareMorphology,
    find_occurrences,
    load_rules,
    sanity_check,
    walk_document,
)


class AiCorrectionEngineSmokeTests(unittest.TestCase):
    def test_public_engine_import_surface_runs_on_fixture(self) -> None:
        fixtures = Path(__file__).parent / "fixtures"
        doc = Document(fixtures / "halyk_finance_sample.docx")
        placeholders = json.loads((fixtures / "halyk_finance_placeholders.json").read_text(encoding="utf-8"))

        text_units = walk_document(doc)
        occurrences = find_occurrences(text_units, placeholders)
        passed, messages = sanity_check(text_units, occurrences)

        self.assertTrue(text_units)
        self.assertTrue(occurrences)
        self.assertTrue(passed, messages)

        engine = CorrectionEngine(
            rules=load_rules("config/ai_governing_phrases.json"),
            morphology=KazakhAwareMorphology(),
            openai_client=None,
        )
        results = engine.correct_document(occurrences)

        self.assertTrue(results)
        self.assertEqual(len(results), len(occurrences))


if __name__ == "__main__":
    unittest.main()
